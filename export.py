import os
import glob
import json
import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

# === Configuration ===
WORLDS_DIR = 'Worlds'
SHEETS_DIR = 'Sheets'


def process_dataframe(records):
    df = pd.DataFrame(records)
    # Numeric conversions
    df['Leve Amount'] = pd.to_numeric(df.get('Leve Amount', 1), errors='coerce').fillna(1)
    df['Leve Gil'] = (
        df.get('Leve Gil', '')
          .astype(str)
          .str.replace(',', '', regex=False)
          .astype(float)
    )

    # Calculate raw material costs
    df['LevePriceNQ'] = df['currentAveragePriceNQ'] * df['Leve Amount']
    df['LevePriceHQ'] = df['currentAveragePriceHQ'] * df['Leve Amount']

    # Masks where item cannot be purchased (price == 0)
    mask_nq_zero = df['currentAveragePriceNQ'] == 0
    mask_hq_zero = df['currentAveragePriceHQ'] == 0

    # Profit calculations
    df['LeveProfitNQ'] = df['Leve Gil'] - df['LevePriceNQ']
    # HQ formula (you had a negative sign on gil ×2—left as you wrote it)
    df['LeveProfitHQ'] = - (df['Leve Gil'] * 2) - df['LevePriceHQ']

    # Zero-price items → set profit to NA and price to zero
    df.loc[mask_nq_zero, ['LevePriceNQ',   'LeveProfitNQ']] = [0, pd.NA]
    df.loc[mask_hq_zero, ['LevePriceHQ',   'LeveProfitHQ']] = [0, pd.NA]

    return df


def main():
    # 1) Verify Worlds dir
    if not os.path.isdir(WORLDS_DIR):
        print(f"[ERROR] Could not find Worlds directory: {WORLDS_DIR}")
        return

    # 2) Ensure Sheets dir exists
    os.makedirs(SHEETS_DIR, exist_ok=True)

    # 3) Loop over each world subfolder
    for world in sorted(os.listdir(WORLDS_DIR)):
        world_path = os.path.join(WORLDS_DIR, world)
        if not os.path.isdir(world_path):
            continue

        json_files = glob.glob(os.path.join(world_path, '*.json'))
        if not json_files:
            print(f"[SKIP] No JSON files in world '{world}'")
            continue

        output_file = os.path.join(SHEETS_DIR, f"{world}_Profits.xlsx")
        print(f"\n▶ Generating: {output_file}")

        # 4) Write each JSON → sheet
        with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
            for file_path in sorted(json_files):
                fname      = os.path.basename(file_path)
                sheet_name = os.path.splitext(fname)[0][:31]  # Excel limit

                try:
                    with open(file_path, 'r', encoding='utf-8') as f:
                        records = json.load(f)
                except Exception as e:
                    print(f"  [ERROR] loading {fname}: {e}")
                    continue

                if not isinstance(records, list) or not records:
                    print(f"  [SKIP] {fname}: no data")
                    continue

                df = process_dataframe(records)
                df.to_excel(writer, sheet_name=sheet_name, index=False)
                print(f"  • Sheet '{sheet_name}' with {len(df)} rows")

        # 5) Re-open with openpyxl to add Tables
        wb = load_workbook(output_file)
        for ws in wb.worksheets:
            ref   = ws.dimensions
            tbl   = Table(displayName=f"Table_{ws.title}", ref=ref)
            style = TableStyleInfo(
                name="TableStyleMedium9",
                showRowStripes=True,
                showColumnStripes=False
            )
            tbl.tableStyleInfo = style
            ws.add_table(tbl)
            print(f"  ✓ Table added on sheet '{ws.title}'")

        wb.save(output_file)
        print(f"✔ Done '{world}' → {output_file}")

    print("\nAll worlds processed.")


if __name__ == '__main__':
    main()
